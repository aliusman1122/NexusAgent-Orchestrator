"""Unified local file loader for A7 Logics Enterprise AI Chatbot.

Scans the local data directory and ingests:
- PDF files (*.pdf) via PyPDFLoader
- Word documents (*.docx) via Docx2txtLoader
- Excel spreadsheets (*.xlsx) via openpyxl, formatting rows into semantic text blocks.

Attaches uniform metadata: {"source": filename, "type": file_extension}
"""

import logging
from pathlib import Path
from typing import List, Optional, Union
import openpyxl
from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader
from langchain_core.documents import Document

from config.settings import settings

logger = logging.getLogger(__name__)


class LocalDocLoader:
    """Unified document parser for local multi-format files."""

    def __init__(self, data_dir: Optional[Union[str, Path]] = None):
        self.data_dir = Path(data_dir or settings.DATA_DIR)
        self.data_dir.mkdir(parents=True, exist_ok=True)

    def load_pdf(self, file_path: Path) -> List[Document]:
        """Load PDF using LangChain PyPDFLoader."""
        logger.info("Loading PDF document: %s", file_path.name)
        try:
            loader = PyPDFLoader(str(file_path))
            raw_docs = loader.load()
            docs = []
            for doc in raw_docs:
                page_num = doc.metadata.get("page", 0)
                metadata = {
                    "source": file_path.name,
                    "source_name": file_path.name,
                    "source_type": "pdf",
                    "type": "pdf",
                    "file_path": str(file_path),
                    "page": page_num + 1,  # 1-indexed for human readability
                }
                docs.append(
                    Document(
                        page_content=doc.page_content.strip(),
                        metadata=metadata,
                    )
                )
            logger.info("Loaded %d pages from %s", len(docs), file_path.name)
            return docs
        except Exception as exc:
            logger.error("Failed to load PDF (%s): %s", file_path.name, exc)
            return []

    def load_docx(self, file_path: Path) -> List[Document]:
        """Load Word document (.docx) using Docx2txtLoader."""
        logger.info("Loading Word document: %s", file_path.name)
        try:
            loader = Docx2txtLoader(str(file_path))
            raw_docs = loader.load()
            docs = []
            for doc in raw_docs:
                content = doc.page_content.strip()
                if not content:
                    continue
                metadata = {
                    "source": file_path.name,
                    "source_name": file_path.name,
                    "source_type": "docx",
                    "type": "docx",
                    "file_path": str(file_path),
                }
                docs.append(Document(page_content=content, metadata=metadata))
            logger.info("Loaded Word document: %s", file_path.name)
            return docs
        except Exception as exc:
            logger.error("Failed to load Word document (%s): %s", file_path.name, exc)
            return []

    def load_xlsx(self, file_path: Path) -> List[Document]:
        """Load Excel spreadsheet (.xlsx) via openpyxl and format rows into semantic text."""
        logger.info("Loading Excel spreadsheet: %s", file_path.name)
        docs = []
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                # First non-empty row as header
                header_row_idx = -1
                headers = []
                for idx, row in enumerate(rows):
                    if any(cell is not None and str(cell).strip() != "" for cell in row):
                        header_row_idx = idx
                        headers = [
                            str(c).strip() if c is not None else f"Column_{i+1}"
                            for i, c in enumerate(row)
                        ]
                        break

                if header_row_idx == -1 or not headers:
                    continue

                sheet_records = []
                for row_num, row in enumerate(rows[header_row_idx + 1 :], start=header_row_idx + 2):
                    # Check if row has any non-empty values
                    row_cells = [
                        (headers[i], str(cell).strip())
                        for i, cell in enumerate(row)
                        if i < len(headers) and cell is not None and str(cell).strip() != ""
                    ]
                    if not row_cells:
                        continue

                    row_description = ", ".join(
                        f"{col}: {val}" for col, val in row_cells
                    )
                    sheet_records.append(
                        f"Row {row_num} -> {row_description}"
                    )

                if sheet_records:
                    semantic_content = (
                        f"Spreadsheet: {file_path.name} | Sheet: {sheet_name}\n"
                        + "\n".join(sheet_records)
                    )
                    metadata = {
                        "source": file_path.name,
                        "source_name": file_path.name,
                        "source_type": "xlsx",
                        "type": "xlsx",
                        "sheet": sheet_name,
                        "file_path": str(file_path),
                        "row_count": len(sheet_records),
                    }
                    docs.append(Document(page_content=semantic_content, metadata=metadata))

            logger.info("Loaded %d sheet(s) from %s", len(docs), file_path.name)
            return docs
        except Exception as exc:
            logger.error("Failed to load Excel spreadsheet (%s): %s", file_path.name, exc)
            return []

    def load_file(self, file_path: Path) -> List[Document]:
        """Load a single file based on its extension."""
        if not file_path.is_file():
            return []
        ext = file_path.suffix.lower()
        if ext == ".pdf":
            return self.load_pdf(file_path)
        elif ext == ".docx":
            return self.load_docx(file_path)
        elif ext == ".xlsx":
            return self.load_xlsx(file_path)
        elif ext in [".txt", ".md", ".csv"]:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read().strip()
                if text:
                    return [
                        Document(
                            page_content=text,
                            metadata={
                                "source": file_path.name,
                                "source_name": file_path.name,
                                "source_type": ext.lstrip("."),
                                "type": ext.lstrip("."),
                                "file_path": str(file_path),
                            },
                        )
                    ]
            except Exception as exc:
                logger.error("Failed to load text document (%s): %s", file_path.name, exc)
                return []
        return []

    def load_all(self) -> List[Document]:
        """Scan data directory for all supported files and return parsed Documents."""
        if not self.data_dir.exists():
            logger.warning("Data directory %s does not exist.", self.data_dir)
            return []

        all_docs: List[Document] = []
        files = list(self.data_dir.rglob("*"))
        logger.info("Found %d file(s) in %s", len(files), self.data_dir)

        for file_path in files:
            if not file_path.is_file():
                continue
            all_docs.extend(self.load_file(file_path))

        logger.info(
            "Completed local doc loading. Total documents parsed: %d", len(all_docs)
        )
        return all_docs


def load_local_documents(data_dir: Optional[Union[str, Path]] = None) -> List[Document]:
    """Convenience helper to load local documents from data directory."""
    loader = LocalDocLoader(data_dir=data_dir)
    return loader.load_all()


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    loaded_docs = load_local_documents()
    print(f"Total documents loaded from data dir: {len(loaded_docs)}")
